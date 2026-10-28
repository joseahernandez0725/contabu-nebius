"""
ContaBu — LOS MODELOS POR ASIENTO, EN EL LIBRO REAL (FASE 11, Auditoría / IA · 22-oct-2026).

¿La elección de modelos hecha con libros sintéticos (`team_matrix.py`, 130 corridas) se sostiene en un libro
REAL, categorizado por un bookkeeper y nunca visto por un modelo? Cada asiento con los mismos candidatos que
`team_matrix` (nano · lightning · super · ultra · vl, sin y con razonar), N corridas por candidato:

- `drafter`  (④, libro real): `field_test.grade_once` — el libro vuelve al "antes" (`--since`), Verificar con la
  evidencia GUARDADA (0 créditos; todos los modelos ven la misma) y el ④ con ese modelo; cobertura y precisión
  contra lo que decidió el bookkeeper, 2026 aparte.
- `personal` (⑤, libro real): la misma llamada que `personal_use.refresh`, con una VERDAD APROXIMADA sacada de las
  categorías del bookkeeper: un comercio cuyas salidas están TODAS en una cuenta personal o en *Shareholder
  distributions* es "personal"; uno con TODAS en gastos del negocio no lo es; los demás no se califican.
- `exam`     (①/⑤, la COPIA `contabu_real_examen`): la revisión completa como `measure.py --audit` (auditar → ②
  con evidencia guardada → ① otra vez → ④), contra la clave de `private/real_exam/key.json`. 🔴 A CIEGAS: el
  script lee la clave para puntuar y solo devuelve conteos; su sha1 se compara antes y después de cada corrida.
- `check`: sin IA y sin créditos — tamaño de la clave tras volver al "antes" y cuánta evidencia guardada hay.
- `table`: las tablas, de `private/field_test/real_matrix.jsonl`.
- `diag` (28-oct): sin IA — el Excel del diagnóstico del ⑤ que guardó `exam --diag` (para que Jose etiquete), y con
  `--score` la precisión del ⑤ con sus etiquetas. Las órdenes, en `no_web_matrix.py`.

🔴 Nada se escribe en las bases: cada corrida es UNA transacción que se deshace (`session.commit` = `flush`; los
agentes no registran o registran dentro de ella). El script comprueba en qué base está (`drafter`/`personal` en
`contabu_real`, `exam` en `contabu_real_examen`). Datos reales: a `private/` con permisos 600; en pantalla, solo
conteos y porcentajes — ningún comercio, descripción ni monto. Ningún número de aquí va al README sin el
consentimiento del dueño del libro.

    U=$(grep ^DATABASE_URL .env.real | cut -d= -f2-)
    DATABASE_URL="$U" .venv/bin/python real_matrix.py check
    DATABASE_URL="$U" .venv/bin/python real_matrix.py drafter --n 3
    DATABASE_URL="$U" .venv/bin/python real_matrix.py personal --n 3
    DATABASE_URL="${U%/*}/contabu_real_examen" .venv/bin/python real_matrix.py exam --n 3
    .venv/bin/python real_matrix.py table
"""

import argparse
import hashlib
import json
import os
import sys
import time
from collections import Counter
from datetime import datetime
from pathlib import Path

from sqlalchemy import text
from sqlmodel import Session, select

HERE = Path(__file__).resolve().parent
OUT = HERE / "private" / "field_test" / "real_matrix.jsonl"
KEY = HERE / "private" / "real_exam" / "key.json"
KEY_SHA1 = HERE / "private" / "real_exam" / "key_frozen_sha1.txt"
COMPANY_ID = 1
SINCE = datetime(2026, 10, 20, 20, 50)          # la prueba de Jose: reloj de la base
MODELS = ("nano", "lightning", "super", "ultra", "vl")
CONFIGS = [(m, r) for m in MODELS for r in (False, True)]
USD_CAP = float(os.getenv("REAL_MATRIX_CAP", "4.00"))   # tope de cada tanda (Nebius); ver `_spent`
# Los exámenes ciegos sembrados (23-oct): base, clave, prefijo del `kind`, cuántos comercios NUEVOS busca el ② por
# revisión y un archivo de evidencia propio (lo de un examen no se mezcla con la evidencia del libro real).
EXAMS = {
    "1": {"db": "contabu_real_examen", "key": "real_exam", "prefix": "exam", "web_limit": 50, "evidence": None},
    "1_2026": {"db": "contabu_real_examen1_2026", "key": "real_exam", "prefix": "exam1_2026", "web_limit": 50,
               "evidence": None},
    # El examen 2 trae filas INSERTADAS (comercios que el ② no conoce): en vivo la primera vez, con tope de 10 comercios
    # nuevos (≤ 30 créditos: 1 básica + 2 avanzada cada uno), y su evidencia en un archivo aparte.
    "2": {"db": "contabu_real_examen2", "key": "real_exam2", "prefix": "exam2", "web_limit": 10,
          "evidence": "web_evidencia_examen2.json"},
}
TAVILY_CAP = 30
SEAT = {"drafter": "drafter", "drafter_note": "drafter", "drafter_note_fix": "drafter", "r3_2026": "drafter",
        "r3_2026_note": "drafter",
        **{f"{e['prefix']}{x}_{k}": seat for e in EXAMS.values() for x in ("", "_nonote")
           for k, seat in (("reviewer", "reviewer"), ("personal", "personal"))}, "personal": "personal", "exam_reviewer": "reviewer",
        "exam_personal": "personal"}
KEY_SIZE = (386, 118)    # la clave del ④ tras volver al "antes" (total, 2026): si cambia, el libro ya no es el de la medición


def _append(row: dict) -> None:
    OUT.parent.mkdir(parents=True, exist_ok=True)
    fd = os.open(OUT, os.O_WRONLY | os.O_CREAT | os.O_APPEND, 0o600)
    os.chmod(OUT, 0o600)
    with os.fdopen(fd, "a", encoding="utf-8") as fh:
        fh.write(json.dumps(row, ensure_ascii=False, default=str) + "\n")


def _read() -> list[dict]:
    return [json.loads(x) for x in open(OUT, encoding="utf-8")] if OUT.exists() else []


SPENT_SINCE = os.getenv("REAL_MATRIX_SINCE", "2026-10-22T20:38")   # tope por tanda: su inicio (reloj del sistema, `ts`)


def _spent() -> float:
    return sum(r.get("cost_usd") or 0.0 for r in _read() if r.get("ts", "") >= SPENT_SINCE)


def _db(session: Session) -> str:
    return session.connection().execute(text("select current_database()")).scalar()


def _require_db(session: Session, name: str) -> None:
    got = _db(session)
    if got != name:
        sys.exit(f"⛔ esta orden corre sobre `{name}` y la conexión es `{got}`")


def _key_sha1() -> str:
    return hashlib.sha1(KEY.read_bytes()).hexdigest()


def _ok(r: dict) -> bool:
    """Una corrida que terminó (un error de red o de la llamada no cuenta: se repite)."""
    return not r.get("error") and str(r.get("outcome", "ok")) in ("ok", "nothing_qualifies")


def _retry(r: dict) -> bool:
    """Solo una falla de RED se repite; un `out_of_tokens` o una respuesta inválida es un resultado del modelo."""
    return "network" in str(r.get("error") or r.get("outcome") or "")


def _done(kind: str, model: str, reasoning: bool) -> int:
    return sum(1 for r in _read() if r["kind"] == kind and r["model"] == model and r["reasoning"] == reasoning
               and not _retry(r))


def _cfg_runs(kind: str, n: int):
    """(modelo, razona) pendientes, en orden: primero la configuración de hoy (el control)."""
    import agent_models
    today = agent_models.DEFAULTS[_seat(kind)]
    order = [today] + [c for c in CONFIGS if c != today]
    for m, r in order:
        for _ in range(max(0, n - _done(kind, m, r))):
            yield m, r


def _seat(kind: str) -> str:
    """El asiento de un `kind` (también los que llevan etiqueta o espacio de nombres)."""
    if kind in SEAT:
        return SEAT[kind]
    if "r3" in kind.split("_"):   # el ④ contra el bookkeeper (`p2_r3`, `p2_r3_note`)
        return "drafter"
    for suffix, seat in (("_reviewer", "reviewer"), ("_personal", "personal"), ("_notes", "notes"),
                         ("_drafter", "drafter"), ("_team", "reviewer")):
        if kind.endswith(suffix):
            return seat
    return "personal"


def _over_cap() -> bool:
    if _spent() >= USD_CAP:
        print(f"⛔ TOPE: ${_spent():.4f} ≥ ${USD_CAP}; me detengo.", flush=True)
        return True
    return False


# ── check ───────────────────────────────────────────────────────────────────────────────────────

def check() -> int:
    import field_test
    import merchant_verifier
    from db import engine
    from models import Company, Transaction
    with Session(engine) as s:
        try:
            _require_db(s, "contabu_real")
            s.commit = s.flush  # type: ignore[method-assign]
            company = s.get(Company, COMPANY_ID)
            back = field_test.rewind(s, company, SINCE)
            key = back["key"]
            txs = s.exec(select(Transaction).where(Transaction.company_id == company.id)).all()
            by_year = Counter("2026" if t.date.year == 2026 else "antes" for t in txs if t.id in key)
            saved = {r["fingerprint"] for r in json.load(open(field_test._oracle_path(company), encoding="utf-8"))}
            flagged, pending, _mem = merchant_verifier.review_candidates(s, company, txs)
            full = {(t.description or "").strip().upper() for t in txs}
            key_fps = {merchant_verifier.query_for(t.description, full).upper() for t in txs if t.id in key}
            searchable = {c.fingerprint for c in merchant_verifier.candidates(txs) if not c.skip}
            print(f"vuelta al 'antes': {back['reverted']} · cambios quitados {back['changes_removed']} · "
                  f"reglas quitadas {back['rules_removed']}")
            print(f"clave: {len(key)} movimientos (2026: {by_year['2026']} · antes: {by_year['antes']})")
            print(f"comercios de la clave que se pueden buscar: {len(key_fps & searchable)} · con evidencia guardada: "
                  f"{len(key_fps & searchable & saved)}")
            print(f"lo que buscaría Verificar: {len(flagged)} marcados + {len(pending)} Pendientes · con evidencia "
                  f"guardada: {sum(c.fingerprint in saved for c in flagged + pending)}")
        finally:
            s.rollback()
    return 0


# ── ④ en el libro real ─────────────────────────────────────────────────────────────────────────

def drafter(n: int, live_first: bool = False, keep_note: int | None = None, kind: str | None = None,
            db: str = "contabu_real") -> int:
    """`db="contabu_real_2026"` (24-oct, regla de Jose: del libro real solo vale 2026): la copia recortada a 2026; su
    clave es solo de 2026 (se exige 0 filas de antes) y los `kind` son `r3_2026` / `r3_2026_note`."""
    import agent_models
    import field_test
    import merchant_verifier
    from db import engine
    from models import Company
    r3 = db != "contabu_real"
    kind = kind or (("r3_2026" if r3 else "drafter") + ("_note" if keep_note is not None else ""))
    keep = frozenset({keep_note}) if keep_note is not None else frozenset()
    first = True
    for m, r in _cfg_runs(kind, n):
        if _over_cap():
            break
        mode, first = ("live" if live_first and first else "saved"), False
        with Session(engine) as s:
            try:
                _require_db(s, db)
                company = s.get(Company, COMPANY_ID)
                agent_models.OVERRIDES["drafter"] = (m, r)
                merchant_verifier.ORACLE = None
                t0 = time.time()
                try:
                    g = field_test.grade_once(s, company, SINCE, mode, keep_note_ids=keep)
                    err = None
                    size = (g["spans"]["total"]["key"], g["spans"]["2026"]["key"])
                    if r3 and g["spans"]["antes"]["key"]:
                        sys.exit(f"⛔ la copia de 2026 trae filas de antes en la clave ({size}); me detengo")
                    if not r3 and size != KEY_SIZE:
                        sys.exit(f"⛔ la clave dio {size}, no {KEY_SIZE}: el libro cambió; me detengo")
                    if keep:   # la nota que se conserva sigue ahí (las demás posteriores sí se borran: 24-oct, nota 2)
                        from models import BookNote
                        left = set(s.exec(select(BookNote.id).where(BookNote.company_id == company.id)).all())
                        if not keep <= left:
                            sys.exit("⛔ se borró la nota que debía quedarse; me detengo")
                        g["notes_left"] = sorted(left)
                except Exception as e:   # un modelo que falla cuenta como corrida fallida, no detiene el resto
                    g, err = None, type(e).__name__
            finally:
                s.rollback()
                agent_models.OVERRIDES.pop("drafter", None)
        row = {"kind": kind, "model": m, "reasoning": r, "keep_note": keep_note, "ts": datetime.now().isoformat(timespec="seconds"),
               "wall_s": round(time.time() - t0, 1), "error": err, "web": mode}
        if g:
            row.update({k: g[k] for k in ("outcome", "calls", "proposals", "rejection_counts", "spans", "cost_usd",
                                          "latency_ms", "completion_tokens", "privacy", "rows", "rejected_amounts")})
            row["notes_left"] = g.get("notes_left")
            row["verify"] = {k: g["verify"][k] for k in ("merchants", "pending", "pending_high_or_medium", "credits")} if g["verify"] else None
        _append(row)
        sp = row.get("spans", {}).get("total", {})
        print(f"{'④+nota' if keep else '④'} {m}{'+r' if r else ''}: {row.get('outcome', err)} · {row.get('proposals', 0)} filas · decide "
              f"{sp.get('decided', 0)}/{sp.get('key', 0)} · acierta {sp.get('right', 0)} · ${row.get('cost_usd', 0):.5f} · "
              f"{(row.get('latency_ms') or 0) / 1000:.1f}s", flush=True)
    return 0


# ── ⑤ en el libro real, con la verdad aproximada ──────────────────────────────────────────────

def _personal_truth(session: Session, company) -> tuple[dict, dict]:
    """llave del comercio → "personal" | "business" (o fuera si se mezcla), y llave → movimientos. Personal = la
    cuenta del libro que dice "personal" o *Shareholder distributions*; negocio = las secciones de gasto."""
    import personal_use
    from models import Transaction
    txs = session.exec(select(Transaction).where(Transaction.company_id == company.id)).all()
    groups: dict[str, list] = {}
    for t in txs:
        if t.is_inflow or not t.category or t.account_type is None:
            continue
        k = personal_use.merchant_key(t)
        if k:
            groups.setdefault(k, []).append(t)
    truth = {}
    for k, g in groups.items():
        kinds = set()
        for t in g:
            cat = t.category.lower()
            if "personal" in cat or "shareholder distribution" in cat:
                kinds.add("personal")
            elif t.account_type.value in personal_use.BUSINESS_SECTIONS:
                kinds.add("business")
            else:
                kinds.add("other")
        if kinds == {"personal"}:
            truth[k] = "personal"
        elif kinds == {"business"}:
            truth[k] = "business"
    return truth, groups


def personal(n: int) -> int:
    import agent_models
    import agents
    import nemotron
    import personal_use
    from db import engine
    from models import Company
    for m, r in _cfg_runs("personal", n):
        if _over_cap():
            break
        with Session(engine) as s:
            try:
                _require_db(s, "contabu_real")
                company = s.get(Company, COMPANY_ID)
                truth, groups = _personal_truth(s, company)
                view = agents.build_view(s, company)
                # Lo que vería el ⑤ si TODO estuviera como gasto del negocio: lo personal primero (son pocos), luego
                # los comercios del negocio por cantidad de movimientos, hasta el mismo tope de producción.
                graded = sorted(truth, key=lambda k: (truth[k] != "personal", -len(groups[k]), k))
                todo = [(k, groups[k]) for k in graded][:personal_use.MAX_MERCHANTS_SENT]
                user, refs = personal_use._user(company, view, todo)
                model_id, reasoning = agent_models.resolve("personal", agent_models.MODEL_IDS[m], r)
                nemotron.SENT.clear()
                t0, err, verdicts, usage = time.time(), None, {}, None
                try:
                    reply = nemotron.chat_json(personal_use._SYSTEM, user, personal_use.JudgmentsOut, model=model_id,
                                               max_tokens=4000, reasoning=reasoning)
                    usage = reply.usage
                    for raw in (reply.data.get("judgments") if isinstance(reply.data, dict) else []) or []:
                        try:
                            j = personal_use.Judgment.model_validate(raw)
                        except Exception:
                            continue
                        if j.ref in refs:
                            verdicts[refs[j.ref][0]] = (j.verdict, j.confidence)
                except nemotron.NemotronError as e:
                    err, usage = e.code, e.usage
                privacy = agents.privacy_scan(list(nemotron.SENT), view.ctx.txs,
                                              __import__("merchant_verifier").protected_names(view.ctx.txs))
            finally:
                s.rollback()
        sent = [k for k, _ in todo]
        alert = {k for k, (v, c) in verdicts.items() if v in personal_use.ALERT_VERDICTS and c in personal_use.ALERT_CONFIDENCE}
        pers = [k for k in sent if truth[k] == "personal"]
        bus = [k for k in sent if truth[k] == "business"]
        row = {"kind": "personal", "model": m, "reasoning": r, "ts": datetime.now().isoformat(timespec="seconds"),
               "wall_s": round(time.time() - t0, 1), "error": err, "sent": len(sent), "answered": len(verdicts),
               "personal_truth": len(pers), "business_truth": len(bus),
               "personal_alerted": sum(k in alert for k in pers), "business_alerted": sum(k in alert for k in bus),
               "verdicts": dict(Counter(v for v, _c in verdicts.values())),
               "cost_usd": (usage.cost_usd if usage else 0.0) or 0.0, "latency_ms": usage.latency_ms if usage else 0,
               "privacy": privacy}
        _append(row)
        print(f"⑤ {m}{'+r' if r else ''}: {err or 'ok'} · contestó {len(verdicts)}/{len(sent)} · personales alertados "
              f"{row['personal_alerted']}/{len(pers)} · del negocio alertados {row['business_alerted']}/{len(bus)} · "
              f"{row['verdicts']} · ${row['cost_usd']:.5f}", flush=True)
    return 0


# ── 28-oct: el diagnóstico del ⑤ con y sin Tavily (`exam --diag`, luego `diag`) ──────────────────────────
# ¿Por qué el ⑤ encuentra 12 de 12 en R2 sin la web y 7 con ella? Por corrida se guarda, en un archivo APARTE (600),
# qué tipo recibió el ⑤ por comercio, qué opinó de cada error de la clave y sus alertas que la clave no tiene, con
# sus filas. 🔴 Trae comercios, montos y la clave: solo lo abre Jose (Excel, `diag`); en pantalla, solo conteos.
DIAG = HERE / "private" / "field_test" / "personal_diag.jsonl"
DIAG_XLSX = HERE / "private" / "field_test" / "personal_diag_examen2.xlsx"
DIAG_ON = False
CONDITION = "with"        # `no_web_matrix.py` lo cambia: "without" (nadie ve el tipo) · "personal_off" (solo el ⑤)
SENT_TO_PERSONAL: dict = {}   # llave del comercio → el tipo que recibió el ⑤ (None = ninguno)
LABELS = ("personal real", "negocio (falsa alarma)", "no sé")


def capture_personal() -> None:
    """Anota lo que de verdad le llega al ⑤ (lo lee de su mensaje, después de cualquier otro parche)."""
    import personal_use
    if getattr(personal_use._user, "_capture", False):
        return
    inner = personal_use._user

    def _user(company, view, todo):
        user, refs = inner(company, view, todo)
        sent = json.loads(user.split("\n\nmerchants: ", 1)[1])
        for m in sent:
            SENT_TO_PERSONAL[refs[m["ref"]][0]] = m.get("verified_business_type")
        return user, refs
    _user._capture = True  # type: ignore[attr-defined]
    personal_use._user = _user


def _diag(session: Session, company, truth: list[dict], rev: dict) -> dict:
    """Lo de UNA corrida, dentro de su transacción (antes del rollback)."""
    import personal_use
    from models import Finding, Transaction
    txs = session.exec(select(Transaction).where(Transaction.company_id == company.id)).all()
    by_id = {t.id: t for t in txs}
    import measure
    by_key: dict[tuple, list[int]] = {}
    for t in txs:
        by_key.setdefault(measure._row_key(t.description, t.date, t.amount, t.is_inflow), []).append(t.id)
    planted = {f"{e['type']}#{i + 1}": {x for r in e["rows"] for x in by_key.get(
        measure._row_key(r["description"], r["date"], r["amount"], r["is_inflow"]), [])} for i, e in enumerate(truth)}
    all_planted = set().union(*planted.values()) if planted else set()
    saved = personal_use.current(session, company)
    tx_row = lambda t: {"id": t.id, "date": str(t.date), "description": t.description, "amount": float(t.amount),  # noqa: E731
                        "category": f"{t.account_type.value}/{t.category}" if t.account_type else None}
    extra = []
    for f in session.exec(select(Finding).where(Finding.company_id == company.id, Finding.type == "personal_use",
                                                Finding.status == "open")).all():
        ids = set(f.transaction_ids or [])
        if ids & all_planted:
            continue
        key = f.key.split(":", 1)[1]
        j = saved.get(key)
        extra.append({"merchant_key": key, "merchant": f.merchant, "type_sent": SENT_TO_PERSONAL.get(key),
                      "verdict": j.verdict if j else None, "confidence": j.confidence if j else None,
                      "reason_es": j.reason_es if j else None,
                      "txs": [tx_row(by_id[i]) for i in sorted(ids) if i in by_id]})
    pers_found = set(planted) - set(rev["personal"]["missed"])
    errors = []
    for i, (k, ids) in enumerate(planted.items()):
        keys = sorted({personal_use.merchant_key(by_id[x]) for x in ids if x in by_id} - {None})
        errors.append({"key": k, "blind": truth[i]["blind"], "merchants": keys,
                       "type_sent": {m: SENT_TO_PERSONAL.get(m) for m in keys},
                       "verdict": {m: f"{saved[m].verdict}/{saved[m].confidence}" if m in saved else None for m in keys},
                       "found_by_personal": k in pers_found, "found_by_team": k in set(rev["found_keys"])})
    return {"alerts": rev["personal"]["alerts"], "true_alerts": rev["personal"]["true_alerts"],
            "extra": extra, "errors": errors}


def _diag_rows(tags: list[str]) -> dict[str, list[dict]]:
    rows = [json.loads(x) for x in open(DIAG, encoding="utf-8")] if DIAG.exists() else []
    return {t: [r for r in rows if r["tag"] == t] for t in tags}


def diag(tags: list[str], score: bool) -> int:
    """Sin IA. Sin `--score`: el Excel para etiquetar (una fila por comercio alertado que la clave no tiene). Con
    `--score`: lee las etiquetas de Jose y da la precisión del ⑤ por corrida. En pantalla, solo conteos."""
    from openpyxl import Workbook, load_workbook
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation
    runs = _diag_rows(tags)
    if score:
        ws = load_workbook(DIAG_XLSX)["Etiquetar"]
        head = [c.value for c in ws[1]]
        label = {r[head.index("merchant_key")]: (r[head.index("etiqueta")] or "").strip().lower()
                 for r in ws.iter_rows(min_row=2, values_only=True) if r[head.index("merchant_key")]}
        for t in tags:
            print(f"{t} ({', '.join(sorted({r['condition'] for r in runs[t]}))}):")
            for r in runs[t]:
                c = Counter(label.get(e["merchant_key"], "") for e in r["extra"])
                real, false = c[LABELS[0]], c[LABELS[1]]
                print(f"  alertas {r['alerts']} · en la clave {r['true_alerts']} · extra {len(r['extra'])}: "
                      f"reales {real}, falsas {false}, sin saber {len(r['extra']) - real - false} → "
                      f"reales en total {r['true_alerts'] + real} de {r['alerts']}")
        return 0
    merchants: dict[str, dict] = {}
    for t in tags:
        for r in runs[t]:
            for e in r["extra"]:
                m = merchants.setdefault(e["merchant_key"], {"e": e, "runs": Counter(), "types": {}, "txs": {}})
                m["runs"][t] += 1
                m["types"].setdefault(t, e["type_sent"])
                m["txs"].update({x["id"]: x for x in e["txs"]})
    wb = Workbook()
    ws = wb.active
    ws.title = "Etiquetar"
    ws.append(["merchant_key", "comercio", "compras", "primera", "última", "monto total", "categoría actual",
               *[f"alertas en {t} (de {len(runs[t])})" for t in tags], *[f"tipo que vio el ⑤ en {t}" for t in tags],
               "razón del ⑤", "etiqueta"])
    for key, m in sorted(merchants.items(), key=lambda kv: -sum(kv[1]["runs"].values())):
        txs = sorted(m["txs"].values(), key=lambda x: x["date"])
        ws.append([key, m["e"]["merchant"], len(txs), txs[0]["date"], txs[-1]["date"],
                   round(sum(abs(x["amount"]) for x in txs), 2), ", ".join(sorted({x["category"] or "" for x in txs})),
                   *[m["runs"][t] for t in tags], *[m["types"].get(t) or "" for t in tags], m["e"]["reason_es"], None])
    dv = DataValidation(type="list", formula1='"' + ",".join(LABELS) + '"', allow_blank=True)
    ws.add_data_validation(dv)
    col = get_column_letter(ws.max_column)
    dv.add(f"{col}2:{col}{max(2, ws.max_row)}")
    ws2 = wb.create_sheet("Compras")
    ws2.append(["merchant_key", "fecha", "descripción", "monto", "categoría actual"])
    for key, m in merchants.items():
        for x in sorted(m["txs"].values(), key=lambda x: x["date"]):
            ws2.append([key, x["date"], x["description"], x["amount"], x["category"]])
    ws3 = wb.create_sheet("Errores de la clave")
    ws3.append(["error", "ciego", "comercio", *[f"{t} #{i + 1}: tipo · opinión · ⑤ · equipo"
                                                  for t in tags for i in range(len(runs[t]))]])
    keys = [(e["key"], e["blind"], m) for r in next((runs[t] for t in tags if runs[t]), [])[:1]
            for e in r["errors"] for m in (e["merchants"] or [None])]
    for k, blind, m in keys:
        cells = []
        for t in tags:
            for r in runs[t]:
                e = next(x for x in r["errors"] if x["key"] == k)
                cells.append(f"{(e['type_sent'].get(m) if m else None) or '—'} · {(e['verdict'].get(m) if m else None) or '—'}"
                             f" · {'sí' if e['found_by_personal'] else 'no'} · {'sí' if e['found_by_team'] else 'no'}")
        ws3.append([k, "sí" if blind else "no", m or "", *cells])
    DIAG_XLSX.parent.mkdir(parents=True, exist_ok=True)
    fd = os.open(DIAG_XLSX, os.O_WRONLY | os.O_CREAT | os.O_TRUNC, 0o600)
    os.chmod(DIAG_XLSX, 0o600)
    with os.fdopen(fd, "wb") as fh:
        wb.save(fh)
    print(f"Excel (600): {DIAG_XLSX.relative_to(HERE)} · {len(merchants)} comercios para etiquetar · "
          + " · ".join(f"{t}: {len(runs[t])} corridas" for t in tags))
    return 0


# ── el examen ciego en la COPIA ─────────────────────────────────────────────────────────────────

def _key_path(exam: str) -> Path:
    return HERE / "private" / EXAMS[exam]["key"] / "key.json"


def _exam_truth(session: Session, company, exam: str = "1") -> list[dict]:
    """La clave → la forma que puntúa `measure.measure_review` (filas por descripción, fecha, monto y dirección),
    con cada fila tomada de la COPIA por su id. `blind` (examen 2): los errores con `blind: false` se puntúan APARTE.
    No se imprime nada de la clave."""
    from models import Transaction
    raw = json.loads(_key_path(exam).read_text(encoding="utf-8"))
    errors = raw if isinstance(raw, list) else next(v for v in raw.values() if isinstance(v, list))
    by_id = {t.id: t for t in session.exec(select(Transaction).where(Transaction.company_id == company.id)).all()}
    out = []
    for e in errors:
        rows = []
        for r in e.get("rows", []):
            t = by_id.get(r.get("id"))
            if t is not None:
                rows.append({"description": t.description, "date": t.date, "amount": t.amount, "is_inflow": t.is_inflow})
        out.append({"type": e.get("type"), "rows": rows, "expected": e.get("expected"),
                    "blind": e.get("blind", True) is not False, "id_rows": len(e.get("rows", []))})
    return out


# ── 24-oct: el ② CON NEMOTRON, sobre resultados CRUDOS guardados (`--web replay`) ─────────────────────────────
# El oráculo guarda tipos YA clasificados (por la tabla): con él, el clasificador nuevo ni corre. Aquí Tavily se
# reemplaza por sus resultados crudos guardados (`private/web_classify/raw.jsonl` + `replay.jsonl`); lo que falte se busca
# UNA vez de verdad y se guarda (600), así las corridas siguientes reusan lo mismo y lo único que cambia es Nemotron.
REPLAY = HERE / "private" / "web_classify" / "replay.jsonl"
REPLAY_STATS: Counter = Counter()


def replay_on() -> None:
    import merchant_verifier as mv
    if getattr(mv._search, "_replay", False):
        return
    real = mv._search
    cache = {}
    for path in (HERE / "private" / "web_classify" / "raw.jsonl", REPLAY):
        if path.exists():
            for x in path.read_text().splitlines():
                if x.strip():
                    r = json.loads(x)
                    cache.setdefault((r["query"], r["level"]), r)

    def fake(query: str, depth: str = "basic"):
        hit = cache.get((query, depth))
        if hit is None:
            results, used, lat = real(query, depth)
            hit = {"set": "replay", "query": query, "level": depth, "credits": used, "latency_ms": lat,
                   "results": [{k: r.get(k) for k in ("title", "content", "url", "score")} for r in results],
                   "ts": time.strftime("%Y-%m-%d %H:%M:%S")}
            REPLAY.parent.mkdir(parents=True, exist_ok=True)
            fd = os.open(REPLAY, os.O_WRONLY | os.O_CREAT | os.O_APPEND, 0o600)
            os.chmod(REPLAY, 0o600)
            with os.fdopen(fd, "a", encoding="utf-8") as fh:
                fh.write(json.dumps(hit, ensure_ascii=False) + "\n")
            cache[(query, depth)] = hit
            REPLAY_STATS["live_credits"] += used
        REPLAY_STATS["searches"] += 1
        REPLAY_STATS["with_results"] += bool(hit["results"])
        return hit["results"], hit["credits"], hit["latency_ms"]
    fake._replay = True  # type: ignore[attr-defined]
    mv._search = fake
    # El control (`REPLAY_CLASSIFIER=table`): la tabla de palabras sobre los MISMOS resultados crudos.
    mv.CLASSIFIER = os.getenv("REPLAY_CLASSIFIER", "nemotron")
    import collections
    import nemotron
    nemotron.SENT = collections.deque(maxlen=5000)   # con el clasificador se envían muchos más mensajes que 50


def privacy_split(sent: list[dict], txs: list, persons: set) -> dict:
    """La privacidad, separada: los agentes de siempre (todo lo enviado) y el clasificador del ②, del que se revisa lo que
    sale DEL LIBRO (la consulta y la región); los resultados son texto de internet (traen precios, que no son montos)."""
    import agents
    mark = "You classify ONE business"
    is_clf = lambda m: m["messages"][0]["content"].startswith(mark)   # noqa: E731
    clf = []
    for m in sent:
        if is_clf(m):
            u = json.loads(m["messages"][1]["content"])
            clf.append({"messages": [{"content": json.dumps({k: v for k, v in u.items() if k != "results"})}]})
    return {"agents": agents.privacy_scan([m for m in sent if not is_clf(m)], txs, persons),
            "classifier_from_book": agents.privacy_scan(clf, txs, persons), "classifier_messages": len(clf)}


def _reclassify_memory(session: Session, company, fingerprints: set) -> int:
    """La memoria de la oficina como la habría construido el ② NUEVO: los mismos comercios que la evidencia guardada,
    clasificados por Nemotron sobre sus resultados crudos (la misma política: básica; avanzada si no salió alta)."""
    import merchant_verifier as mv
    from models import Transaction
    txs = session.exec(select(Transaction).where(Transaction.company_id == company.id)).all()
    region, cache, n = mv.book_region(txs), mv.cached(session, company), 0
    for c in mv.candidates(txs):
        if c.fingerprint in fingerprints and c.fingerprint not in cache and not c.skip:
            mv._lookup(session, company, c, cache, record=False, region=region)
            n += 1
    session.flush()
    return n


def _evidence_paths(company, exam: str) -> list[Path]:
    import field_test
    extra = EXAMS[exam]["evidence"]
    return [field_test._oracle_path(company)] + ([field_test.OUT_DIR / extra] if extra else [])


def _clf_stats(session: Session, books: list) -> dict:
    """Las llamadas del clasificador del ② en ESTA revisión (las de la memoria reconstruida no se registran)."""
    from models import AgentCall
    rows = session.exec(select(AgentCall).where(AgentCall.company_id.in_([b.id for b in books]),
                                                AgentCall.agent == "web_classifier")).all()
    per = Counter(r.company_id for r in rows)
    model_calls = Counter(r.company_id for r in session.exec(select(AgentCall).where(
        AgentCall.company_id.in_([b.id for b in books]), AgentCall.agent != "verifier")).all())
    return {"clf_calls": len(rows), "clf_fallback": sum(r.outcome != "ok" for r in rows),
            "clf_usd": round(sum(r.cost_usd or 0 for r in rows), 6), "clf_max_per_book": max(per.values(), default=0),
            "model_calls_max_per_book": max(model_calls.values(), default=0), "model_calls_all_books": sum(model_calls.values()),
            "replay_searches": REPLAY_STATS["searches"], "replay_with_results": REPLAY_STATS["with_results"],
            "replay_live_credits": REPLAY_STATS["live_credits"]}


def exam_once(session: Session, company, *, oracle: bool, exam: str = "1", drop_note: bool = False,
              replay: bool = False) -> dict:
    """UNA revisión completa sobre la copia, como `measure.py --audit --web oracle`: los veredictos y las opiniones
    del ⑤ de antes se borran DENTRO de la transacción, para que cada modelo juzgue desde cero."""
    import agents
    import field_test
    import measure
    import merchant_verifier
    import nemotron
    from models import BookNote, Finding, MerchantJudgment, Transaction
    session.commit = session.flush  # type: ignore[method-assign]
    if drop_note:   # el equipo de hoy SIN las notas del libro (dentro de la transacción)
        for n in session.exec(select(BookNote).where(BookNote.company_id == company.id)).all():
            session.delete(n)
    for j in session.exec(select(MerchantJudgment).where(MerchantJudgment.company_id == company.id)).all():
        session.delete(j)
    for f in session.exec(select(Finding).where(Finding.company_id == company.id)).all():
        session.delete(f)
    session.flush()
    # La evidencia guardada entra como MEMORIA de la oficina (dentro de la transacción): en vivo solo se busca lo
    # que falta; con `oracle`, lo que falta sale "no encontrado" sin red (0 créditos).
    from models import MerchantLookup
    have = set(merchant_verifier.cached(session, company))
    saved_rows = [r for path in _evidence_paths(company, exam) if path.exists()
                  for r in json.load(open(path, encoding="utf-8"))]
    reclassified = _reclassify_memory(session, company, {r["fingerprint"] for r in saved_rows}) if replay else 0
    for r in ([] if replay else saved_rows):
        if r["fingerprint"] not in have:
            have.add(r["fingerprint"])
            session.add(MerchantLookup(firm_id=company.firm_id, company_id=None if company.firm_id else company.id,
                                       fingerprint=r["fingerprint"], query=r["fingerprint"],
                                       status="ok" if r.get("results_count") else "not_found",
                                       business_type=r.get("business_type"), confidence=r.get("confidence"),
                                       source_url=r.get("source_url"), results_count=r.get("results_count") or 0,
                                       credits=0, search_depth=r.get("search_depth") or "basic"))
    session.flush()
    merchant_verifier.ORACLE = {} if (oracle and not replay) else None
    nemotron.SENT.clear()
    REPLAY_STATS.clear()
    SENT_TO_PERSONAL.clear()
    agents.CITATION_REJECTED.clear()
    truth = _exam_truth(session, company, exam)
    before = set(merchant_verifier.cached(session, company))
    # El flujo de `measure.web_flow` (auditar → ② → ① otra vez) con el tope de búsquedas de ESTE examen.
    agents.review(session, company, lang="en")
    vres = merchant_verifier.verify_flagged(session, company, limit=EXAMS[exam]["web_limit"])
    if vres.finding_ids:
        agents.review(session, company, lang="en", finding_ids=vres.finding_ids)
    web = {"searched": vres.searched, "credits": vres.credits}
    if not oracle and not replay:   # la corrida en vivo: su evidencia se SUMA a la guardada, para que las demás la reusen (0 créditos)
        path = _evidence_paths(company, exam)[-1]
        saved = {r["fingerprint"]: r for r in json.load(open(path, encoding="utf-8"))} if path.exists() else {}
        for fp, m in merchant_verifier.cached(session, company).items():
            if fp not in before:
                saved[fp] = {"fingerprint": m.fingerprint, "business_type": m.business_type, "confidence": m.confidence,
                             "source_url": m.source_url, "results_count": m.results_count, "search_depth": m.search_depth}
        with field_test._private_open(path) as fh:
            json.dump(list(saved.values()), fh, ensure_ascii=False, indent=2)
    pend = agents.draft(session, company, record=False)
    rev = measure.measure_review(session, company, truth, draft_rows=[p.row for p in pend.proposals],
                                 already_reviewed=True)
    costs = measure.run_costs(session, [company])
    txs = session.exec(select(Transaction).where(Transaction.company_id == company.id)).all()
    from models import AgentCall
    lat = sum(c.latency_ms for c in session.exec(select(AgentCall).where(
        AgentCall.company_id == company.id, AgentCall.agent != "verifier", AgentCall.created_at >= SINCE)).all())
    by_type = Counter(k.split("#")[0] for k in rev["found_keys"])
    planted_types = Counter(e["type"] for e in truth)
    keys = [f"{e['type']}#{i + 1}" for i, e in enumerate(truth)]
    blind = {k for k, e in zip(keys, truth) if e["blind"]}
    code_keys = set(keys) - set(rev["auditor"]["missed"])
    team_keys = set(rev["found_keys"])
    pers_keys = set(keys) - set(rev["personal"]["missed"])
    note_keys = set(keys) - set(rev["notes"]["missed"])
    privacy = (privacy_split if replay else agents.privacy_scan)(list(nemotron.SENT), txs, merchant_verifier.protected_names(txs))
    a, rv, team = rev["auditor"], rev["reviewer"], rev["team"]
    return {"planted": rev["planted"], "unmatched": len(rev["unmatched"]),
            "blind_planted": len(blind), "open_planted": len(keys) - len(blind),
            "code_blind": len(code_keys & blind), "code_open": len(code_keys - blind),
            "agents_blind": len(team_keys & blind), "agents_open": len(team_keys - blind),
            "personal_blind": len(pers_keys & blind), "personal_open": len(pers_keys - blind),
            "personal_false_alerts": rev["personal"]["alerts"] - rev["personal"]["true_alerts"],
            "citation_rejected": dict(Counter(agents.CITATION_REJECTED)),
            "notes_blind": len(note_keys & blind), "notes_open": len(note_keys - blind),
            "notes_false_alerts": rev["notes"]["alerts"] - rev["notes"]["true_alerts"],
            "rows_in_key": sum(e["id_rows"] for e in truth), "rows_found": sum(len(e["rows"]) for e in truth),
            "code_found": a["found"], "code_alerts": a["alerts"], "code_on_planted": a["true_alerts"],
            "agents_found": team["found"], "agents_alerts": team["alerts"], "agents_on_planted": team["true_alerts"],
            "reviewer_found": rv["found"], "personal_found": rev["personal"]["found"],
            "personal_alerts": rev["personal"]["alerts"], "drafter_found": rev["drafter"]["found"],
            "dismissed_true": len(rev["dismissed_true"]), "judged": rev["judged"],
            "found_by_type": dict(by_type), "planted_by_type": dict(planted_types),
            "searched": web["searched"], "credits_if_live": web["credits"], **_clf_stats(session, [company]),
            "memory_reclassified": reclassified,
            "cost_usd": costs["nebius_usd"] + ((pend.usage.cost_usd or 0.0) if pend.usage else 0.0),
            "nebius_calls": costs["nebius_calls"] + pend.calls,
            "latency_ms": lat + (pend.usage.latency_ms if pend.usage else 0), "privacy": privacy,
            **({"_diag": _diag(session, company, truth, rev)} if DIAG_ON else {})}


def exam(n: int, live_first: bool, exam_id: str = "1", drop_note: bool = False, tag: str | None = None,
         ns: str | None = None, seat: str = "team", replay: bool = False) -> int:
    import agent_models
    from db import engine
    from models import Company
    cfg = EXAMS[exam_id]
    key_path = _key_path(exam_id)
    sha0 = (key_path.parent / "key_frozen_sha1.txt").read_text().split()[0]
    sha = lambda: hashlib.sha1(key_path.read_bytes()).hexdigest()   # noqa: E731
    prefix = cfg["prefix"] + ("_nonote" if drop_note else "") + (f"_{tag}" if tag else "")
    if ns:   # 24-oct: un espacio de nombres (p. ej. los prompts reescritos) y UN asiento, o el equipo de hoy
        prefix = f"{cfg['prefix']}_{ns}" + ("_nonote" if drop_note else "")
        kinds = [(f"{prefix}_team", "reviewer")] if seat == "team" else [(f"{prefix}_{seat}", seat)]
        tag = tag or ("team" if seat == "team" else None)
    else:
        kinds = [(f"{prefix}_reviewer", "reviewer"), (f"{prefix}_personal", "personal")]
    plan = []
    for kind, seat in kinds:
        runs = list(_cfg_runs(kind, n))
        if drop_note or tag:   # sin nota, o una medición con etiqueta: solo el equipo de hoy
            runs = [(m, r) for m, r in runs if (m, r) == agent_models.DEFAULTS[seat]][:n]
            if kind.endswith("_personal"):
                runs = []   # el equipo de hoy es UNO: con el asiento del ① basta
        plan += [(kind, seat, m, r) for m, r in runs]
    first = True
    for kind, seat, m, r in plan:
        if _over_cap():
            break
        if sha() != sha0:
            sys.exit("⛔ la clave cambió (sha1): me detengo")
        with Session(engine) as s:
            try:
                _require_db(s, cfg["db"])
                company = s.get(Company, COMPANY_ID)
                agent_models.OVERRIDES[seat] = (m, r)
                t0 = time.time()
                try:
                    out, err = exam_once(s, company, oracle=not (live_first and first), exam=exam_id,
                                         drop_note=drop_note, replay=replay), None
                except Exception as e:
                    out, err = None, type(e).__name__
            finally:
                s.rollback()
                agent_models.OVERRIDES.pop(seat, None)
        diag_row = (out or {}).pop("_diag", None)
        if sha() != sha0:
            sys.exit("⛔ la clave cambió (sha1) durante la corrida: me detengo")
        if out and out.get("credits_if_live", 0) > TAVILY_CAP and live_first and first:
            print(f"⚠️ Tavily: {out['credits_if_live']} créditos (> {TAVILY_CAP})", flush=True)
        row = {"kind": kind, "model": m, "reasoning": r, "ts": datetime.now().isoformat(timespec="seconds"),
               "exam": exam_id, "drop_note": drop_note,
               "wall_s": round(time.time() - t0, 1), "error": err,
               "web": "replay" if replay else "live" if (live_first and first) else "saved", "web_condition": CONDITION,
               **(out or {"cost_usd": 0.0})}
        first = False
        _append(row)
        if diag_row is not None:
            DIAG.parent.mkdir(parents=True, exist_ok=True)
            fd = os.open(DIAG, os.O_WRONLY | os.O_CREAT | os.O_APPEND, 0o600)
            os.chmod(DIAG, 0o600)
            with os.fdopen(fd, "a", encoding="utf-8") as fh:
                fh.write(json.dumps({"tag": tag, "kind": kind, "condition": CONDITION, "ts": row["ts"], **diag_row},
                                    ensure_ascii=False, default=str) + "\n")
        print(f"{kind} {m}{'+r' if r else ''}: {err or 'ok'} · ciegos: código {row.get('code_blind')} · con agentes "
              f"{row.get('agents_blind')} de {row.get('blind_planted')} · no ciegos {row.get('agents_open')}/{row.get('open_planted')} · "
              f"⑤ {row.get('personal_blind')} (+{row.get('personal_false_alerts')} sin clasificar) · "
              f"filas de la clave {row.get('rows_found')}/{row.get('rows_in_key')} · créditos {row.get('credits_if_live')} · "
              f"código solo {row.get('code_found')}/{row.get('planted')} · "
              f"con agentes {row.get('agents_found')}/{row.get('planted')} · ⑤ {row.get('personal_found')} · "
              f"① despachó {row.get('dismissed_true')} verdaderos · ${row.get('cost_usd', 0):.5f} · "
              f"{(row.get('latency_ms') or 0) / 1000:.1f}s", flush=True)
    return 0


# ── el demo: D1 y D2 en UNA revisión (24-oct) ───────────────────────────────────────────────────

def demo_once(replay: bool = False) -> dict:
    """Los 3 libros del demo sembrados (en una transacción que se deshace) y UNA revisión completa, como
    `team_matrix.run_once` (memoria de la oficina vacía, ② por oráculo, 0 créditos). Se puntúa DOS veces con la misma
    revisión: D1 (el primer examen, escondiendo las filas del ciego, como siempre) y D2 (el ciego, con el primer examen
    como neutral). El ciego no cambia lo que hacen los agentes, solo la clave: así cada corrida cuesta la mitad."""
    import demo_data
    import demo_web_evidence
    import formula
    import measure
    import merchant_verifier
    import nemotron
    import agents
    from db import engine
    from models import AgentCall, Firm, MerchantLookup, Transaction, User
    merchant_verifier.ORACLE = None if replay else {r["fingerprint"]: r for r in demo_web_evidence.load()}
    nemotron.SENT.clear()
    REPLAY_STATS.clear()
    agents.CITATION_REJECTED.clear()
    out = Counter()
    with Session(engine) as s:
        try:
            s.commit = s.flush  # type: ignore[method-assign]
            owner = User(firebase_uid=f"__p2_{time.time_ns()}__", email="p2@test.local"); s.add(owner); s.flush()
            firm = Firm(name="P2 DEMO", owner_user_id=owner.id); s.add(firm); s.flush()
            books = demo_data.seed_demo_books(s, firm, owner)
            for m in s.exec(select(MerchantLookup).where(MerchantLookup.firm_id == firm.id)).all():
                s.delete(m)
            s.flush()
            all_txs = []
            for b in books:
                tp, te1 = measure._truth(b.name, blind=False)
                _tp, te2 = measure._truth(b.name, blind=True)
                txs = s.exec(select(Transaction).where(Transaction.company_id == b.id)).all()
                all_txs += txs
                measure.web_flow(s, b)
                pend = measure.measure_pending(s, b, tp, record=True)
                nc = measure.note_case(formula.BookContext(s, b), pend["rows"]) or {}
                r1 = measure.measure_review(s, b, te1, draft_rows=pend["rows"], hide=frozenset(measure._blind_ids(txs)),
                                            already_reviewed=True)
                r2 = measure.measure_review(s, b, te2, draft_rows=pend["rows"],
                                            other_exam=frozenset(measure._exam_ids(txs, "DEMO_TRUTH_ERRORS")),
                                            already_reviewed=True)
                for tag, r in (("d1", r1), ("d2", r2)):
                    out[f"{tag}_planted"] += r["team"]["planted"]
                    out[f"{tag}_found"] += r["team"]["found"]
                    out[f"{tag}_code"] += r["auditor"]["found"]
                    out[f"{tag}_alerts"] += r["team"]["alerts"]
                    out[f"{tag}_true_alerts"] += r["team"]["true_alerts"]
                    out[f"{tag}_dismissed_true"] += len(r["dismissed_true"])
                    out[f"{tag}_personal_found"] += r["personal"]["found"]
                    out[f"{tag}_personal_false"] += r["personal"]["alerts"] - r["personal"]["true_alerts"]
                    out[f"{tag}_notes_found"] += r["notes"]["found"]
                    out[f"{tag}_notes_false"] += r["notes"]["alerts"] - r["notes"]["true_alerts"]
                out["pending_known"] += pend["known"]
                out["pending_decided"] += pend["decided"]
                out["pending_right"] += pend["right"]
                out["vacation_expected"] += nc.get("expected", 0)
                out["vacation_hit"] += nc.get("hit", 0)
                out["vacation_extra"] += nc.get("extra", 0)
            calls = s.exec(select(AgentCall).where(AgentCall.company_id.in_([b.id for b in books]))).all()
            out["cost_usd"] = sum(c.cost_usd or 0 for c in calls if c.agent != "verifier")
            out["latency_ms"] = sum(c.latency_ms for c in calls if c.agent != "verifier")
            out["out_of_tokens"] = sum(1 for c in calls if c.outcome == "out_of_tokens")
            privacy = (privacy_split if replay else agents.privacy_scan)(list(nemotron.SENT), all_txs,
                                                                         merchant_verifier.protected_names(all_txs))
            out.update(_clf_stats(s, books))
        finally:
            s.rollback()
    return {**dict(out), "citation_rejected": dict(Counter(agents.CITATION_REJECTED)), "privacy": privacy}


def demo(n: int, seat: str, ns: str, replay: bool = False) -> int:
    """`seat="team"`: el equipo de hoy × n. Si no, ESE asiento varía (los 10 candidatos × n) y los demás quedan en el de
    hoy. Las filas: `kind = f"{ns}_demo_{seat}"`."""
    import agent_models
    kind = f"{ns}_demo_{seat}"
    configs = [agent_models.DEFAULTS["reviewer"]] if seat == "team" else None
    today = agent_models.DEFAULTS[seat] if seat != "team" else None
    order = [(m, r) for m, r in ([today] + [c for c in CONFIGS if c != today])] if seat != "team" else configs
    for m, r in order:
        for _ in range(max(0, n - _done(kind, m, r))):
            if _over_cap():
                return 1
            agent_models.OVERRIDES.clear()
            if seat != "team":
                agent_models.OVERRIDES[seat] = (m, r)
            t0 = time.time()
            try:
                out, err = demo_once(replay), None
            except Exception as e:
                out, err = {"cost_usd": 0.0}, type(e).__name__
            finally:
                agent_models.OVERRIDES.clear()
            row = {"kind": kind, "model": m, "reasoning": r, "ts": datetime.now().isoformat(timespec="seconds"),
                   "wall_s": round(time.time() - t0, 1), "error": err, "web": "replay" if replay else "oracle",
                   "web_condition": CONDITION, **out}
            _append(row)
            print(f"{kind} {m}{'+r' if r else ''}: {err or 'ok'} · D1 {out.get('d1_found')}/{out.get('d1_planted')} · "
                  f"D2 {out.get('d2_found')}/{out.get('d2_planted')} (código {out.get('d2_code')}) · ④ {out.get('pending_right')}/"
                  f"{out.get('pending_decided')} de {out.get('pending_known')} · vacaciones {out.get('vacation_hit')}/"
                  f"{out.get('vacation_expected')} · ⑥ {out.get('d1_notes_found')}+{out.get('d2_notes_found')} · "
                  f"${out.get('cost_usd', 0):.4f}", flush=True)
    return 0


# ── tablas ──────────────────────────────────────────────────────────────────────────────────────

def _fmt(vals, f="{}"):
    return " · ".join(f.format(v) for v in vals) if vals else "—"


def table() -> int:
    rows = _read()
    print(f"Gastado en Nebius: ${_spent():.4f} (tope ${USD_CAP})\n")
    known = ["drafter", "drafter_note", "drafter_note_fix", "r3_2026", "r3_2026_note", "personal", "exam_reviewer",
             "exam_personal"]
    kinds = known + sorted({r["kind"] for r in rows} - set(known))
    for kind in kinds:
        mine = [r for r in rows if r["kind"] == kind]
        if not mine:
            continue
        print(f"== {kind}")
        for m, rz in CONFIGS:
            rs = [r for r in mine if r["model"] == m and r["reasoning"] == rz and not _retry(r)]
            if not rs:
                continue
            ok = [r for r in rs if _ok(r)]
            cost = sum(r.get("cost_usd") or 0 for r in rs) / len(rs)
            lat = sum((r.get("latency_ms") or 0) for r in ok) / max(1, len(ok)) / 1000
            name = f"{m}{' razona' if rz else ''}"
            if kind in ("drafter", "drafter_note", "drafter_note_fix", "r3_2026", "r3_2026_note"):
                cov = [f"{r['spans']['total']['decided']}/{r['spans']['total']['key']}" for r in ok]
                prec = [f"{r['spans']['total']['right']}/{r['spans']['total']['decided']}" for r in ok]
                c26 = [f"{r['spans']['2026']['decided']}/{r['spans']['2026']['key']}·{r['spans']['2026']['right']}" for r in ok]
                rej = Counter()
                for r in ok:
                    rej.update(r.get("rejection_counts") or {})
                print(f"  {name:16} decide {_fmt(cov)} · acierta {_fmt(prec)} · 2026 (decide·acierta) {_fmt(c26)} · "
                      f"filas {_fmt([r['proposals'] for r in ok])} · rechazos {dict(rej)} · fallas {len(rs) - len(ok)} · "
                      f"${cost:.5f} · {lat:.1f}s")
            elif kind == "personal":
                print(f"  {name:16} personales alertados {_fmt([f'{r['personal_alerted']}/{r['personal_truth']}' for r in ok])} · "
                      f"del negocio alertados {_fmt([f'{r['business_alerted']}/{r['business_truth']}' for r in ok])} · "
                      f"contestó {_fmt([f'{r['answered']}/{r['sent']}' for r in ok])} · fallas {len(rs) - len(ok)} · "
                      f"${cost:.5f} · {lat:.1f}s")
            elif "blind_planted" in (ok[0] if ok else {}):
                b = ok[0]["blind_planted"]
                print(f"  {name:16} CIEGOS (de {b}): código {_fmt([r['code_blind'] for r in ok])} · con agentes "
                      f"{_fmt([r['agents_blind'] for r in ok])} · ⑤ {_fmt([r['personal_blind'] for r in ok])} "
                      f"(⑤ sin clasificar {_fmt([r['personal_false_alerts'] for r in ok])}) · no ciegos (de {ok[0]['open_planted']}) "
                      f"{_fmt([r['agents_open'] for r in ok])} · ① despachó verdaderos {_fmt([r['dismissed_true'] for r in ok])} · "
                      f"alertas sembradas/sin clasificar {_fmt([f'{r['agents_on_planted']}/{r['agents_alerts'] - r['agents_on_planted']}' for r in ok])} · "
                      f"fallas {len(rs) - len(ok)} · ${cost:.5f} · {lat:.1f}s")
            else:
                print(f"  {name:16} código solo {_fmt([r['code_found'] for r in ok])} · con agentes "
                      f"{_fmt([r['agents_found'] for r in ok])} de {ok[0]['planted'] if ok else '?'} · ⑤ "
                      f"{_fmt([r['personal_found'] for r in ok])} · ① despachó verdaderos {_fmt([r['dismissed_true'] for r in ok])} · "
                      f"alertas (sobre sembradas / sin clasificar) {_fmt([f'{r['agents_on_planted']}/{r['agents_alerts'] - r['agents_on_planted']}' for r in ok])} · "
                      f"fallas {len(rs) - len(ok)} · ${cost:.5f} · {lat:.1f}s")
        if kind.endswith(("_reviewer", "_personal")) and any("blind_planted" in r for r in mine):
            _adoption(mine, kind)
        print()
    return 0


def _adoption(mine: list[dict], kind: str) -> None:
    """La regla de siempre: un candidato gana si su PEOR corrida supera la MEJOR del modelo de hoy (errores ciegos)."""
    import agent_models
    today = agent_models.DEFAULTS[_seat(kind)]
    score = lambda m, r: [x["agents_blind"] for x in mine if x["model"] == m and x["reasoning"] == r and _ok(x)]   # noqa: E731
    best_today = max(score(*today), default=None)
    if best_today is None:
        return
    wins = [f"{m}{' razona' if r else ''}" for m, r in CONFIGS if (m, r) != today and score(m, r) and min(score(m, r)) > best_today]
    print(f"  → regla de adopción (peor del candidato > mejor de hoy = {best_today}): {', '.join(wins) or 'ninguno gana'}")


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("what", choices=("check", "drafter", "personal", "exam", "demo", "table", "diag"))
    ap.add_argument("--seat", default="team", choices=("team", "reviewer", "personal", "drafter", "notes"),
                    help="demo/exam: el asiento que varía (team = el equipo de hoy)")
    ap.add_argument("--ns", help="exam/demo: espacio de nombres de los kinds (p. ej. p2 = los prompts reescritos)")
    ap.add_argument("--n", type=int, default=3, help="corridas por configuración")
    ap.add_argument("--exam", choices=tuple(EXAMS), default="1", help="exam: cuál examen sembrado (base, clave, kind)")
    ap.add_argument("--tag", help="exam: medición con etiqueta (solo el equipo de hoy, kind aparte)")
    ap.add_argument("--drop-note", action="store_true", help="exam: el equipo de hoy con las notas borradas (en la transacción)")
    ap.add_argument("--db", default="contabu_real", help="drafter: la base (contabu_real_2026 = la copia recortada a 2026)")
    ap.add_argument("--keep-note", type=int, help="drafter: conservar esta nota al volver al 'antes' (kind drafter_note)")
    ap.add_argument("--live-first", action="store_true",
                    help="drafter/exam: la PRIMERA corrida busca en Tavily de verdad y guarda la evidencia; las demás la reusan")
    ap.add_argument("--web", choices=("saved", "replay"), default="saved",
                    help="demo/exam: replay = el ② con Nemotron sobre los resultados CRUDOS guardados (24-oct)")
    ap.add_argument("--diag", action="store_true",
                    help="exam: guarda además el diagnóstico del ⑤ (comercios y la clave: archivo aparte, 600)")
    ap.add_argument("--tags", nargs="+", help="diag: las etiquetas a comparar (p. ej. w28 nw28 p28)")
    ap.add_argument("--score", action="store_true", help="diag: la precisión del ⑤ con las etiquetas de Jose")
    args = ap.parse_args()
    if args.what == "diag":
        if not args.tags:
            sys.exit("⛔ diag necesita --tags")
        return diag(args.tags, args.score)
    if args.diag:
        if args.what != "exam" or not args.tag:
            sys.exit("⛔ --diag es solo para `exam` con --tag")
        global DIAG_ON
        DIAG_ON = True
        capture_personal()
    if args.web == "replay":
        replay_on()
    if args.what == "table":
        return table()
    if args.what == "check":
        return check()
    if args.what == "drafter":
        return drafter(args.n, args.live_first, args.keep_note, db=args.db)
    if args.what == "personal":
        return personal(args.n)
    if args.what == "demo":
        return demo(args.n, args.seat, args.ns or "p2", replay=args.web == "replay")
    return exam(args.n, args.live_first, args.exam, args.drop_note, args.tag, args.ns, args.seat,
                replay=args.web == "replay")


if __name__ == "__main__":
    sys.exit(main())
